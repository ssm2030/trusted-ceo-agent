from __future__ import annotations

import csv
import hashlib
import os
import secrets
import stat
import unicodedata
import xml.etree.ElementTree as ElementTree
import zipfile
from collections.abc import Iterable, Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from trusted_ceo_agent.canonical import strict_loads
from trusted_ceo_agent.errors import ContractError
from trusted_ceo_agent.filesystem import ensure_within
from trusted_ceo_agent.intake.adapters.xlsx_preflight import preflight_xlsx


_CONTENT_TYPES: Mapping[str, frozenset[str]] = {
    ".csv": frozenset({"text/csv"}),
    ".json": frozenset({"application/json"}),
    ".xlsx": frozenset({
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }),
}
_FORBIDDEN_INTERMEDIATE_SUFFIXES = frozenset({
    ".7z", ".bat", ".cmd", ".com", ".csv", ".dll", ".exe", ".gz",
    ".jar", ".js", ".json", ".msi", ".ps1", ".py", ".rar", ".scr",
    ".sh", ".tar", ".vbs", ".xlsm", ".xlsx", ".zip",
})
_WINDOWS_RESERVED_NAMES = frozenset({
    "con", "prn", "aux", "nul",
    *(f"com{index}" for index in range(1, 10)),
    *(f"lpt{index}" for index in range(1, 10)),
})


@dataclass(frozen=True, slots=True)
class UploadLimits:
    max_file_bytes: int = 25 * 1024 * 1024
    max_files: int = 64
    max_total_bytes: int = 250 * 1024 * 1024
    chunk_bytes: int = 1024 * 1024

    def __post_init__(self) -> None:
        if min(
            self.max_file_bytes,
            self.max_files,
            self.max_total_bytes,
            self.chunk_bytes,
        ) <= 0:
            raise ValueError("upload limits must be positive")
        if self.max_file_bytes > self.max_total_bytes:
            raise ValueError("per-file upload limit cannot exceed total limit")


@dataclass(frozen=True, slots=True)
class IncomingUpload:
    filename: str
    content_type: str
    chunks: Iterable[bytes] = field(repr=False)

    @classmethod
    def from_bytes(
        cls,
        filename: str,
        content_type: str,
        payload: bytes,
    ) -> IncomingUpload:
        return cls(filename=filename, content_type=content_type, chunks=(payload,))


@dataclass(frozen=True, slots=True)
class StagedUpload:
    opaque_token: str
    filename: str
    content_type: str
    size: int
    sha256: str
    private_path: Path = field(repr=False)

    def public_metadata(self) -> dict[str, Any]:
        return {
            "opaque_token": self.opaque_token,
            "filename": self.filename,
            "content_type": self.content_type,
            "size": self.size,
            "sha256": self.sha256,
        }


def _validated_name(filename: str) -> tuple[str, str]:
    if not isinstance(filename, str) or not filename or len(filename) > 240:
        raise ContractError("upload filename is invalid")
    filename = unicodedata.normalize("NFC", filename)
    if (
        filename.startswith(".")
        or filename.endswith((".", " "))
        or any(part.endswith(" ") for part in filename.split("."))
        or ":" in filename
        or any(separator in filename for separator in ("/", "\\"))
    ):
        raise ContractError("hidden or path-like upload filename is forbidden")
    path = Path(filename)
    suffixes = [suffix.casefold() for suffix in path.suffixes]
    if not suffixes or suffixes[-1] not in _CONTENT_TYPES:
        raise ContractError("upload extension must end in .csv, .json, or .xlsx")
    if any(suffix in _FORBIDDEN_INTERMEDIATE_SUFFIXES for suffix in suffixes[:-1]):
        raise ContractError("executable, archive, or deceptive double extension is forbidden")
    stem = path.stem
    if (
        not stem
        or stem.casefold().split(".", 1)[0] in _WINDOWS_RESERVED_NAMES
        or any(unicodedata.category(character) == "Cc" for character in filename)
    ):
        raise ContractError("upload filename is invalid")
    return filename, suffixes[-1]


def _validate_content_type(extension: str, content_type: str) -> None:
    normalized = content_type.split(";", 1)[0].strip().casefold()
    if normalized not in _CONTENT_TYPES[extension]:
        raise ContractError(
            f"upload content type does not match {extension}: {content_type}"
        )


def _validate_csv(path: Path) -> None:
    with path.open("rb") as binary:
        prefix = binary.read(4)
    if prefix.startswith((b"PK\x03\x04", b"MZ")):
        raise ContractError("CSV magic bytes are invalid")
    try:
        found_value = False
        with path.open("r", encoding="utf-8-sig", newline="") as text:
            for row in csv.reader(text):
                if any("\x00" in cell for cell in row):
                    raise ContractError("CSV NUL bytes are forbidden")
                if any(cell.strip() for cell in row):
                    found_value = True
    except (UnicodeDecodeError, csv.Error) as error:
        raise ContractError("CSV must be valid UTF-8 tabular text") from error
    if not found_value:
        raise ContractError("CSV must contain at least one non-empty cell")


def _validate_json(payload: bytes) -> None:
    if payload.startswith((b"PK\x03\x04", b"MZ")):
        raise ContractError("JSON magic bytes are invalid")
    try:
        strict_loads(payload)
    except (UnicodeDecodeError, ValueError) as error:
        raise ContractError("JSON upload is invalid") from error


def _validate_xlsx_structure(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            names = {item.filename.replace("\\", "/") for item in archive.infolist()}
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise ContractError("ZIP container is not an XLSX workbook")
            try:
                content_types = ElementTree.fromstring(
                    archive.read("[Content_Types].xml")
                )
            except ElementTree.ParseError as error:
                raise ContractError("XLSX content types XML is invalid") from error
            for item in content_types.iter():
                content_type = item.attrib.get("ContentType", "").casefold()
                if "macroenabled" in content_type or "vba" in content_type:
                    raise ContractError("macro-enabled XLSX content is forbidden")
            for name in sorted(names):
                if not name.casefold().endswith(".rels"):
                    continue
                try:
                    root = ElementTree.fromstring(archive.read(name))
                except ElementTree.ParseError as error:
                    raise ContractError("XLSX relationship XML is invalid") from error
                for relationship in root.iter():
                    if relationship.attrib.get("TargetMode", "").casefold() == "external":
                        raise ContractError("external XLSX relationships are forbidden")
    except zipfile.BadZipFile as error:
        raise ContractError("invalid XLSX ZIP container") from error
    workbook = None
    try:
        with path.open("rb") as handle:
            workbook = load_workbook(
                handle,
                read_only=True,
                data_only=True,
                keep_links=False,
            )
            if not workbook.sheetnames:
                raise ContractError("XLSX workbook must contain a worksheet")
    except (
        ElementTree.ParseError,
        InvalidFileException,
        KeyError,
        OSError,
        ValueError,
    ) as error:
        raise ContractError("XLSX workbook structure is invalid") from error
    finally:
        if workbook is not None:
            workbook.close()


class UploadPolicy:
    def __init__(
        self,
        service_root: Path,
        *,
        limits: UploadLimits | None = None,
    ) -> None:
        service_root.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.service_root = ensure_within(service_root, service_root)
        self.staging_root = ensure_within(
            self.service_root,
            self.service_root / "staging",
        )
        self.staging_root.mkdir(mode=0o700, exist_ok=True)
        self.limits = limits or UploadLimits()

    def stage_batch(
        self,
        uploads: Iterable[IncomingUpload],
        *,
        existing_file_count: int,
        existing_total_bytes: int,
    ) -> tuple[StagedUpload, ...]:
        items = tuple(uploads)
        if not items:
            raise ContractError("at least one upload is required")
        if (
            isinstance(existing_file_count, bool)
            or not isinstance(existing_file_count, int)
            or existing_file_count < 0
            or isinstance(existing_total_bytes, bool)
            or not isinstance(existing_total_bytes, int)
            or existing_total_bytes < 0
        ):
            raise ContractError("existing upload usage is invalid")
        if existing_file_count + len(items) > self.limits.max_files:
            raise ContractError("upload file count exceeds the run limit")
        if existing_total_bytes > self.limits.max_total_bytes:
            raise ContractError("existing uploads exceed the run byte limit")
        staged: list[StagedUpload] = []
        total_size = existing_total_bytes
        try:
            for upload in items:
                item = self._stage_one(upload, already_staged_bytes=total_size)
                staged.append(item)
                total_size += item.size
            return tuple(staged)
        except Exception:
            for item in staged:
                self._unlink_if_private(item.private_path)
            raise

    def _stage_one(
        self,
        upload: IncomingUpload,
        *,
        already_staged_bytes: int,
    ) -> StagedUpload:
        filename, extension = _validated_name(upload.filename)
        _validate_content_type(extension, upload.content_type)
        token = "upload_" + secrets.token_urlsafe(24)
        token_root = ensure_within(
            self.staging_root,
            self.staging_root / token,
        )
        token_root.mkdir(mode=0o700)
        destination = ensure_within(
            token_root,
            token_root / filename,
        )
        digest = hashlib.sha256()
        size = 0
        try:
            with destination.open("xb") as handle:
                for chunk in upload.chunks:
                    if not isinstance(chunk, bytes):
                        raise ContractError("upload stream chunks must be bytes")
                    if not chunk:
                        continue
                    size += len(chunk)
                    if size > self.limits.max_file_bytes:
                        raise ContractError("upload file exceeds 25 MiB")
                    if already_staged_bytes + size > self.limits.max_total_bytes:
                        raise ContractError("upload batch exceeds 250 MiB")
                    digest.update(chunk)
                    handle.write(chunk)
                handle.flush()
                os.fsync(handle.fileno())
            if size == 0:
                raise ContractError("empty upload is forbidden")
            item = StagedUpload(
                opaque_token=token,
                filename=filename,
                content_type=upload.content_type,
                size=size,
                sha256=digest.hexdigest(),
                private_path=destination,
            )
            self.validate_staged(item)
            if extension == ".csv":
                _validate_csv(destination)
            elif extension == ".json":
                _validate_json(destination.read_bytes())
            else:
                with destination.open("rb") as handle:
                    magic = handle.read(4)
                if magic != b"PK\x03\x04":
                    raise ContractError("XLSX magic bytes are invalid")
                preflight_xlsx(destination)
                _validate_xlsx_structure(destination)
            return item
        except Exception:
            self._unlink_if_private(destination)
            raise

    def validate_staged(self, item: StagedUpload) -> Path:
        safe = ensure_within(self.staging_root, item.private_path)
        if (
            safe.name != item.filename
            or safe.parent.name != item.opaque_token
            or safe.parent.parent != self.staging_root
        ):
            raise ContractError("staged upload path is invalid")
        if not safe.is_file():
            raise ContractError("staged upload is missing")
        before_path = safe.stat()
        if getattr(before_path, "st_nlink", 1) != 1:
            raise ValueError("hard-linked staging file is forbidden")
        digest = hashlib.sha256()
        size = 0
        with safe.open("rb") as handle:
            before_handle = os.fstat(handle.fileno())
            while True:
                chunk = handle.read(self.limits.chunk_bytes)
                if not chunk:
                    break
                digest.update(chunk)
                size += len(chunk)
            after_handle = os.fstat(handle.fileno())
        after_path = safe.stat()
        identities = {
            (
                int(getattr(details, "st_dev", 0)),
                int(getattr(details, "st_ino", 0)),
                int(details.st_size),
                int(details.st_mtime_ns),
                int(getattr(details, "st_nlink", 1)),
            )
            for details in (before_path, before_handle, after_handle, after_path)
        }
        if len(identities) != 1 or stat.S_ISLNK(after_path.st_mode):
            raise ContractError("staged upload changed while being read")
        if size != item.size or digest.hexdigest() != item.sha256:
            raise ContractError("staged upload changed after validation")
        return safe

    def discard(self, item: StagedUpload) -> None:
        self.validate_staged(item)
        self._unlink_if_private(item.private_path)

    def _unlink_if_private(self, path: Path) -> None:
        try:
            safe = ensure_within(self.staging_root, path)
        except ValueError:
            return
        if safe.parent.parent == self.staging_root:
            safe.unlink(missing_ok=True)
            try:
                safe.parent.rmdir()
            except OSError:
                pass
