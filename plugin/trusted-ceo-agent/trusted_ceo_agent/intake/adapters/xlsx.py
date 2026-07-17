from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from trusted_ceo_agent.errors import ContractError
from trusted_ceo_agent.intake.adapters.xlsx_preflight import XlsxLimits, preflight_xlsx
from trusted_ceo_agent.intake.models import ParsedDataset, ParsedRecord
from trusted_ceo_agent.intake.quality import quality_issue


def _cell_value(value: Any) -> Any:
    if isinstance(value, float):
        return Decimal(str(value))
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


class XlsxAdapter:
    def __init__(self, *, limits: XlsxLimits | None = None) -> None:
        self.limits = limits or XlsxLimits()

    def parse(self, path: Path, source_id: str) -> ParsedDataset:
        preflight_xlsx(path, self.limits)
        handle = None
        try:
            # Immutable Source blobs are content-addressed and intentionally have
            # no filename extension. A binary stream avoids openpyxl's filename
            # suffix heuristic while preserving the ZIP preflight above.
            handle = path.open("rb")
            workbook = load_workbook(handle, read_only=True, data_only=False, keep_links=False)
        except Exception as error:
            if handle is not None:
                handle.close()
            raise ContractError(f"cannot parse XLSX: {path.name}") from error
        records: list[ParsedRecord] = []
        issues: list[dict[str, Any]] = []
        all_fields: set[str] = set()
        logical_index = 0
        try:
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows()
                try:
                    header_cells = next(rows)
                except StopIteration:
                    continue
                fields = tuple("" if cell.value is None else str(cell.value) for cell in header_cells)
                if not fields or any(not field for field in fields) or len(fields) != len(set(fields)):
                    raise ContractError(f"XLSX sheet {sheet.title} requires unique non-empty headers")
                all_fields.update(fields)
                for row_number, cells in enumerate(rows, start=2):
                    if all(cell.value is None for cell in cells):
                        continue
                    logical_index += 1
                    values: dict[str, Any] = {}
                    for field, cell in zip(fields, cells):
                        if cell.data_type == "f":
                            values[field] = None
                            issues.append(quality_issue(
                                issue_code="untrusted_formula_value",
                                severity="blocking",
                                reason_code="formula_value_requires_confirmation",
                                source_ref={"source_id": source_id, "sheet": sheet.title, "cell": cell.coordinate},
                                affected_field=field,
                                raw_value=cell.value,
                                suggested_resolution="Recalculate and export trusted values or confirm at Data Gate",
                            ))
                        else:
                            values[field] = _cell_value(cell.value)
                    end_column = get_column_letter(max(1, len(fields)))
                    records.append(ParsedRecord(
                        logical_index,
                        "xlsx_cells",
                        {"sheet": sheet.title, "range": f"A{row_number}:{end_column}{row_number}"},
                        values,
                    ))
        finally:
            workbook.close()
            if handle is not None:
                handle.close()
        return ParsedDataset(
            source_id,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            tuple(sorted(all_fields)),
            tuple(records),
            {"sheet_count": len(workbook.sheetnames)},
            issues,
        )
